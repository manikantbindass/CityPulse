"""Editable 33-slide CityPulse presentation from frozen report/data evidence."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import zipfile
import textwrap
import runpy

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
import qrcode

ROOT=Path(__file__).resolve().parents[3]
SOURCE=Path(__file__).resolve().parent
OUTPUT=ROOT/'docs/presentations/CityPulse-Project-Presentation.pptx'
DATA=json.loads((SOURCE/'slides.json').read_text(encoding='utf-8'))
BG='090D16';SURFACE='111B2B';INK='E8EDF5';MUTED='A6B4C9';CYAN='67E8F9';VIOLET='B7A5FF';INDIGO='879BFF';RULE='30445E'
FONT='Arial'
HASHES={}
route_edge=runpy.run_path(str(ROOT/'docs/reports/source/routing.py'))['route_edge']


def asset(name):
    path=(ROOT/name).resolve()
    if not path.is_relative_to(ROOT) or not path.is_file():
        raise ValueError(f'Invalid presentation source: {name}')
    content=path.read_bytes()
    if path.suffix not in {'.png','.jpg','.pdf','.pptx'}:content=content.replace(b'\r\n',b'\n')
    HASHES[path.relative_to(ROOT).as_posix()]=hashlib.sha256(content).hexdigest()
    return path


def read(name):return json.loads(asset(name).read_text(encoding='utf-8'))


def color(value):return RGBColor.from_string(value)


def text(slide,value,x,y,w,h,size=22,fill=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for index,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if index==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.space_after=Pt(7);p.line_spacing=1.12
    return shape


def line(slide,x1,y1,x2,y2,fill=RULE,width=1):
    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=color(fill);shape.line.width=Pt(width)
    return shape


def picture(slide,name,x,y,w,h):
    path=asset(name)
    with Image.open(path) as image:iw,ih=image.size
    scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
    return slide.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))


def points(slide,items,x=.7,y=2.0,w=11.7,size=25,gap=1.2):
    for index,item in enumerate(items):
        text(slide,f'{index+1:02d}',x,y+index*gap,.6,.5,18,CYAN,True)
        text(slide,item,x+.85,y+index*gap,w-.85,gap-.15,size)


def frame(prs,item,index):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(BG)
    if item['kind']!='cover':
        text(slide,item['title'],.65,.48,12.0,1.0,33,INK,True)
        line(slide,.65,1.48,12.68,1.48,CYAN,1.1)
        text(slide,f'CITYPULSE  |  SYNTHETIC DEVELOPMENT / IMPLEMENTATION EVIDENCE',.65,7.05,11.3,.25,10,MUTED)
        text(slide,f'{index:02d} / 33',11.78,7.05,.9,.25,10,MUTED)
    transition=OxmlElement('p:transition');transition.set('spd','fast');transition.append(OxmlElement('p:fade'));slide._element.append(transition)
    sources=[]
    for name in item['sources']:
        asset(name)
        if name.startswith(('docs/reports/source/','docs/presentations/source/')):
            sources.append(f"{name} (companion editable source in this delivery)")
        else:
            sources.append(f"{name}\n{DATA['repository']}/blob/{DATA['snapshot']}/{name}")
    slide.notes_slide.notes_text_frame.text=(item['notes']+'\n\nApplication snapshot: '+DATA['snapshot']+'\nEvidence date: '+DATA['date']+'\nSources:\n'+'\n'.join(sources))
    return slide


def graph(slide,key):
    specs=read('docs/reports/source/diagrams.json');spec=specs[key]
    nodes=spec['nodes'];positions={}
    # Report labels are preserved as editable native text, with a wider slide layout.
    x0,y0=4.35,1.9;w,h=2.55,.78;dx,dy=2.78,1.13
    for i,node in enumerate(nodes):positions[node[0]]=(x0+(i%3)*dx,y0+(i//3)*dy,w,h)
    # Orthogonal connectors run before nodes so no line crosses a label.
    for start,end,label in spec['edges']:
        route=route_edge(start,end,positions,.14)
        for a,b in zip(route,route[1:]):connector=line(slide,*a,*b,CYAN,1)
        end_marker=OxmlElement('a:tailEnd');end_marker.set('type','triangle');connector.line._get_or_add_ln().append(end_marker)
    for ident,title,description in nodes:
        x,y,w,h=positions[ident]
        box=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        box.fill.solid();box.fill.fore_color.rgb=color(SURFACE);box.line.color.rgb=color(RULE);box.line.width=Pt(.8)
        display_title=title.replace('SourceRegistry','Source\nRegistry').replace('AttestationRegistry','Attestation\nRegistry').replace('/', ' / ')
        display='\n'.join(part for row in display_title.splitlines() for part in textwrap.wrap(row,width=23,break_long_words=False,break_on_hyphens=False))
        label=text(slide,display,x+.12,y+.10,w-.24,h-.12,15,INK,True)
        label.name='Diagram label: '+title


def native_chart(slide,categories,series,x,y,w,h,kind=XL_CHART_TYPE.COLUMN_CLUSTERED,legend=True):
    data=CategoryChartData();data.categories=categories
    for name,values in series:data.add_series(name,values)
    graphic=slide.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data)
    chart=graphic.chart;chart.chart_style=10;chart.has_legend=legend
    if legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.font.name=FONT;chart.legend.font.size=Pt(14)
    chart.font.name=FONT;chart.font.size=Pt(14);chart.font.color.rgb=color(MUTED)
    chart.category_axis.tick_labels.font.name=FONT;chart.category_axis.tick_labels.font.size=Pt(14);chart.category_axis.tick_labels.font.color.rgb=color(MUTED)
    chart.value_axis.tick_labels.font.name=FONT;chart.value_axis.tick_labels.font.size=Pt(14);chart.value_axis.tick_labels.font.color.rgb=color(MUTED)
    chart.value_axis.has_major_gridlines=True
    chart.value_axis.major_gridlines.format.line.color.rgb=color(RULE)
    for index,series_obj in enumerate(chart.series):
        series_obj.format.fill.solid();series_obj.format.fill.fore_color.rgb=color([CYAN,VIOLET,INDIGO][index%3])
        series_obj.format.line.color.rgb=color([CYAN,VIOLET,INDIGO][index%3])
    # Set chart and plot backgrounds explicitly instead of relying on a white theme.
    for target in [chart._chartSpace,chart._chartSpace.chart.plotArea]:
        sppr=target.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
        if sppr is None:sppr=OxmlElement('c:spPr');target.append(sppr)
        fill=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val',BG);fill.append(rgb);sppr.append(fill)
        border=OxmlElement('a:ln');border.append(OxmlElement('a:noFill'));sppr.append(border)
    return chart


def process(slide,item):
    descriptions={
        'solution':[('Observe','Compatible evidence'),('Forecast','Named methods'),('Simulate','Explicit constraints'),('Verify','Digest and provenance')],
        'backend':[('Authenticate','Token and membership'),('Authorize','Tenant and scopes'),('Transact','Validate and persist'),('Explain','Safe typed response')],
        'journey':[('Ingest','Validate and replay'),('Investigate','Compare and forecast'),('Decide','Human-reviewed scenario'),('Preserve','Save, export, verify')],
    }
    text(slide,item['points'][0],.7,1.85,11.8,.7,25)
    for i,(label,description) in enumerate(descriptions[item['visual']]):
        x=.7+i*3.1
        text(slide,f'{i+1:02d}',x,3.0,2.5,.8,42,CYAN,True)
        text(slide,label,x,4.0,2.5,.6,27,INK,True)
        text(slide,description,x,4.85,2.5,1.1,19,MUTED)
        if i<3:line(slide,x+2.6,3.55,x+2.95,3.55,VIOLET,2)
    text(slide,'Human review remains the decision boundary',.7,6.2,11.7,.5,18,VIOLET)


def build():
    assert len(DATA['slides'])==33
    asset('docs/presentations/source/slides.json')
    asset('docs/presentations/source/build.py')
    asset('docs/reports/source/routing.py')
    report_builder=runpy.run_path(str(asset('docs/reports/source/build.py')))
    if DATA['snapshot'] != report_builder['CONFIG']['snapshot']:
        raise ValueError('Presentation and report must describe the same application snapshot')
    report_builder['verify_application_snapshot']()
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='CityPulse - AI-Powered Urban Intelligence & Blockchain Analytics'
    prs.core_properties.subject='Technical project presentation; synthetic demonstration'
    prs.core_properties.author='CityPulse repository owner: manikantbindass'
    prs.core_properties.keywords='CityPulse; urban analytics; synthetic data; implementation evidence'
    model=read('docs/reports/model-evaluation.json')
    load=read('docs/reports/operations/k6-local-ai.json')
    for index,item in enumerate(DATA['slides'],1):
        slide=frame(prs,item,index);kind=item['kind']
        if kind=='cover':
            text(slide,'CITYPULSE',.7,1.35,12,1.2,58,CYAN,True)
            text(slide,'AI-Powered Urban Intelligence\n& Blockchain Analytics',.73,2.9,11.9,1.4,30)
            line(slide,.75,4.75,4.0,4.75,VIOLET,3)
            text(slide,'Predict Demand. Optimize Supply.\nBuild Smarter Cities.',.75,5.15,10.6,1.1,23)
            text(slide,DATA['repository'],.75,6.75,11.5,.3,14,MUTED)
        elif kind=='screenshot':
            points(slide,item['points'],.65,2.0,3.4,19,1.35)
            picture(slide,item['visual'],4.35,1.8,8.35,4.9)
        elif kind=='diagram':
            points(slide,item['points'],.65,2.0,3.35,18,1.45)
            graph(slide,item['visual'])
        elif kind=='process':process(slide,item)
        elif kind=='model-chart':
            labels=[key.replace('amsterdam','AMS').replace('bengaluru','BLR').replace('central','C').replace('north','N').replace('groceries','G').replace('mobility','M') for key in model['segments']]
            values=[record['candidates'][record['selectedModel']]['test']['mae'] for record in model['segments'].values()]
            native_chart(slide,labels,[('Selected-model test MAE',values)],.7,2.05,8.6,4.35,XL_CHART_TYPE.BAR_CLUSTERED,False)
            text(slide,'Synthetic\nheld-out\none-step MAE',9.65,2.3,2.9,1.6,26,CYAN,True)
            text(slide,'Compare within segment.\nNo real-city accuracy claim.',9.65,4.3,2.7,1.25,18,MUTED)
            text(slide,'AMS/BLR = Amsterdam/Bengaluru; C/N = central/north; G/M = groceries/mobility',.75,6.55,11.8,.3,13,MUTED)
        elif kind=='map-forecast':
            picture(slide,'docs/presentations/source/evidence/explorer.png',.7,1.9,5.8,4.6)
            demo=read('docs/presentations/source/evidence/forecast-demo.json')['response'];predictions=demo['predictions']
            categories=[row['time'][5:10] for row in predictions]
            series=[('Forecast',[r['demand'] for r in predictions])]
            if all(r.get('lower') is not None and r.get('upper') is not None for r in predictions):
                series.extend([('Lower',[r['lower'] for r in predictions]),('Upper',[r['upper'] for r in predictions])])
            native_chart(slide,categories,series,6.8,2.0,5.85,4.15,XL_CHART_TYPE.LINE,True)
            text(slide,'Synthetic daily demand units, Jul–Aug 2025; residual intervals are diagnostic',.75,6.5,11.9,.4,16,MUTED)
        elif kind=='test-metrics':
            rows=[('29','Java'),('48','AI'),('13','Web unit'),('18','Desktop/mobile'),('6','Ingestion'),('4','Solidity')]
            for i,(value,label) in enumerate(rows):
                x=.8+(i%3)*4.2;y=1.95+(i//3)*2.05
                text(slide,value,x,y,3.6,.9,50,CYAN if i%2==0 else VIOLET,True)
                text(slide,label,x,y+1,3.5,.5,23)
            text(slide,'Revision-linked CI evidence; two mocked Terraform plans; no coverage percentage invented',.8,6.5,11.9,.35,14,MUTED)
        elif kind=='load-chart':
            metric=load['measured']['metrics']['http_req_duration']['values']
            native_chart(slide,['Average','Median','p95','Maximum'],[('Milliseconds',[metric['avg'],metric['med'],metric['p(95)'],metric['max']])],.7,2.0,8.2,4.3,XL_CHART_TYPE.COLUMN_CLUSTERED,False)
            text(slide,'21 requests\n0% HTTP failures',9.3,2.4,3.25,1.6,29,CYAN,True)
            text(slide,'20 seconds\n1 arrival/second\nDirect loopback AI',9.3,4.3,3.25,1.5,20,MUTED)
            text(slide,'Synthetic smoke workload; excludes web/API/DB latency and is not production capacity evidence',.75,6.55,11.9,.35,14,MUTED)
        elif kind=='qa':
            text(slide,'Questions?',.7,2.2,8,1.2,54,CYAN,True)
            text(slide,'Data • models • security • deployment',.75,3.75,8.1,.6,25)
            text(slide,DATA['repository'],.75,5.1,8,.8,20,VIOLET)
            qr=SOURCE/'evidence/repository-qr.png';qrcode.make(DATA['repository']).save(qr)
            picture(slide,qr.relative_to(ROOT).as_posix(),9.6,2.1,2.75,2.75)
            text(slide,'Repository and reproducible evidence',9.35,5.15,3.4,.8,17,MUTED)
        else:
            points(slide,item['points'],.7,2.0,11.8,25,1.25)
            if kind=='roadmap':text(slide,'PLANNED / NOT DEPLOYED',.75,6.5,11.8,.35,16,VIOLET,True)
    OUTPUT.parent.mkdir(exist_ok=True);prs.save(OUTPUT)
    # Reopen, not just write a ZIP with a .pptx suffix.
    reopened=Presentation(OUTPUT)
    if len(reopened.slides)!=33:raise ValueError('Unexpected slide count')
    manifest={'snapshot':DATA['snapshot'],'slideCount':33,'size':'16:9 widescreen','font':'Arial (design choice, installed system font; no font asset committed)','sourceHashes':dict(sorted(HASHES.items())),'outputSha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),'notes':'Every slide has source-linked speaker notes; charts/architecture labels are editable native objects; fade transition only','limitations':'Synthetic evidence; no production, public-chain or real-city accuracy claim'}
    (SOURCE/'evidence.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print(f'Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes), reopened {len(reopened.slides)} slides')


if __name__=='__main__':build()

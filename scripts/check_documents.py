"""Validate committed documents and their normalized source/evidence fingerprints."""
import hashlib
import json
from pathlib import Path
import zipfile

from pypdf import PdfReader
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    content = path.read_bytes()
    if path.suffix not in {'.png', '.jpg', '.pdf', '.pptx'}:
        content = content.replace(b'\r\n', b'\n')
    return hashlib.sha256(content).hexdigest()


def verify_manifest(folder, output):
    manifest = json.loads((ROOT / folder / 'source/evidence.json').read_text(encoding='utf-8'))
    assert digest(ROOT / output) == manifest['outputSha256'], f'Changed output: {output}'
    for name, expected in manifest['sourceHashes'].items():
        path = (ROOT / name).resolve()
        assert path.is_relative_to(ROOT), f'Unsafe source: {name}'
        assert path.is_file() and digest(path) == expected, f'Stale source: {name}; regenerate documents'


def check():
    report = 'docs/reports/CityPulse-Technical-Report.pdf'
    deck = 'docs/presentations/CityPulse-Project-Presentation.pptx'
    verify_manifest('docs/reports', report)
    verify_manifest('docs/presentations', deck)
    reader = PdfReader(ROOT / report, strict=True)
    assert 50 <= len(reader.pages) <= 80, 'Unexpected report length'
    assert len(reader.outline) >= 30, 'Missing chapter bookmarks'
    assert '/JavaScript' not in reader.trailer['/Root'].get('/Names', {}), 'Active PDF content'
    body = '\n'.join(page.extract_text() for page in reader.pages)
    for phrase in ['Contents', 'Synthetic', 'PostGIS', 'Blockchain', 'References']:
        assert phrase.lower() in body.lower(), f'Missing report content: {phrase}'
    embedded = False
    for page in reader.pages:
        for ref in page['/Resources'].get('/Font', {}).values():
            font = ref.get_object()
            descriptor = font.get('/FontDescriptor')
            if descriptor and descriptor.get_object().get('/FontFile2'):
                embedded = True
    assert embedded, 'No embedded report font'
    presentation = Presentation(ROOT / deck)
    assert len(presentation.slides) == 33
    assert abs(presentation.slide_width / presentation.slide_height - 16/9) < .001
    charts = 0
    for index, slide in enumerate(presentation.slides, 1):
        assert slide.has_notes_slide and 'Sources:' in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, f'Negative position: slide {index}'
            assert shape.left + shape.width <= presentation.slide_width + 10, f'Off-slide width: {index}'
            assert shape.top + shape.height <= presentation.slide_height + 10, f'Off-slide height: {index}'
            charts += bool(shape.has_chart)
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert (run.font.name or paragraph.font.name) == 'Arial', f'Unexpected font: slide {index}'
    assert charts == 3, 'Measured charts missing'
    with zipfile.ZipFile(ROOT / deck) as package:
        assert package.testzip() is None
        assert len([p for p in package.namelist() if p.startswith('ppt/embeddings/') and p.endswith('.xlsx')]) == 3
    print(f'Documents passed: {len(reader.pages)} PDF pages; 33 slides; 33 notes; 3 editable charts; source hashes verified')


if __name__ == '__main__':
    check()
